"""Prepare a brand research workbook and export Agent-friendly JSON."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


BRIEF_FIELDS = {
    "任务名称": "task_name",
    "品牌名称": "brand_name",
    "品牌别名": "brand_aliases",
    "行业": "industry",
    "商品名称": "product_name",
    "商品类别": "product_category",
    "上级品类": "parent_category",
    "核心卖点": "selling_points",
    "成分关键词": "ingredient_keywords",
    "目标人群": "target_audience",
    "竞品品牌": "competitor_brands",
    "竞品商品": "competitor_products",
    "必查关键词": "required_keywords",
    "排除关键词": "excluded_keywords",
    "研究目标": "research_goal",
    "重点研究问题": "research_questions",
    "首选时间范围": "preferred_time_range",
    "备注": "notes",
}

REQUIRED_FIELDS = {
    "task_name",
    "brand_name",
    "industry",
    "product_category",
    "research_goal",
    "preferred_time_range",
}

MULTI_VALUE_FIELDS = {
    "brand_aliases",
    "selling_points",
    "ingredient_keywords",
    "target_audience",
    "competitor_brands",
    "competitor_products",
    "required_keywords",
    "excluded_keywords",
    "research_questions",
}

REVIEW_FIELDS = {
    "任务名称": "task_name",
    "关键词": "keyword",
    "关键词类型": "keyword_type",
    "来源字段": "source_field",
    "是否必查": "required",
    "是否采用": "approved",
    "采用/删除原因": "reason",
    "优先级": "priority",
    "备注": "notes",
}

KEYWORD_RULES = (
    ("brand_name", "品牌名称", "品牌", True, "yes", 1, "核心品牌词"),
    ("brand_aliases", "品牌别名", "品牌别名", True, "yes", 1, "品牌检索别名"),
    ("product_name", "商品名称", "商品", True, "yes", 1, "核心商品词"),
    ("product_category", "商品类别", "商品类别", True, "yes", 1, "核心品类词"),
    ("required_keywords", "必查关键词", "必查关键词", True, "yes", 1, "用户指定必查"),
    ("parent_category", "上级品类", "上级品类", False, "", 2, "扩大品类覆盖"),
    ("competitor_brands", "竞品品牌", "竞品品牌", False, "", 2, "竞品对比候选"),
    ("competitor_products", "竞品商品", "竞品商品", False, "", 2, "竞品商品候选"),
    ("selling_points", "核心卖点", "核心卖点", False, "", 3, "卖点趋势候选"),
    ("ingredient_keywords", "成分关键词", "成分", False, "", 3, "成分趋势候选"),
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="生成关键词审核表、校验品牌研究输入，并导出 Agent JSON。"
    )
    parser.add_argument("workbook", type=Path, help="research_input.xlsx 路径")
    parser.add_argument(
        "--output",
        type=Path,
        help="输出 JSON 路径；默认写到 Excel 同目录下。",
    )
    parser.add_argument(
        "--no-sync-keywords",
        action="store_true",
        help="只校验并导出 JSON，不更新 keyword_review。",
    )
    return parser.parse_args()


def normalize_cell(value: Any) -> str:
    return str(value or "").strip()


def split_multi_value(value: str) -> list[str]:
    return [part.strip() for part in value.replace("，", "、").split("、") if part.strip()]


def keyword_values(value: Any) -> list[str]:
    if isinstance(value, list):
        return [normalize_cell(item) for item in value if normalize_cell(item)]
    return split_multi_value(normalize_cell(value))


def generate_keyword_candidates(brief: dict[str, Any]) -> list[dict[str, Any]]:
    excluded = {item.casefold() for item in keyword_values(brief.get("excluded_keywords"))}
    seen: set[str] = set()
    candidates: list[dict[str, Any]] = []

    for field, source_field, keyword_type, required, approved, priority, reason in KEYWORD_RULES:
        for keyword in keyword_values(brief.get(field)):
            normalized = keyword.casefold()
            if normalized in excluded or normalized in seen:
                continue
            seen.add(normalized)
            candidates.append(
                {
                    "task_name": brief["task_name"],
                    "keyword": keyword,
                    "keyword_type": keyword_type,
                    "source_field": source_field,
                    "required": "yes" if required else "no",
                    "approved": approved,
                    "reason": reason,
                    "priority": priority,
                    "notes": "程序生成",
                }
            )
    return candidates


def find_header_row(sheet, expected_header: str, max_rows: int = 20) -> int:
    for row_index in range(1, min(sheet.max_row, max_rows) + 1):
        if any(normalize_cell(cell.value) == expected_header for cell in sheet[row_index]):
            return row_index
    raise ValueError(f"工作表 {sheet.title!r} 中找不到表头 {expected_header!r}")


def read_briefs(workbook_path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "research_brief" not in workbook.sheetnames:
        raise ValueError("缺少 research_brief 工作表")

    sheet = workbook["research_brief"]
    header_row = find_header_row(sheet, "任务名称")
    headers = [normalize_cell(cell.value) for cell in sheet[header_row]]
    header_map = {
        index: BRIEF_FIELDS[header]
        for index, header in enumerate(headers)
        if header in BRIEF_FIELDS
    }

    missing_headers = set(BRIEF_FIELDS) - set(headers)
    if missing_headers:
        raise ValueError(f"输入表缺少字段：{', '.join(sorted(missing_headers))}")

    briefs: list[dict[str, Any]] = []
    errors: list[str] = []
    for row_index, row in enumerate(
        sheet.iter_rows(min_row=header_row + 1, values_only=True),
        start=header_row + 1,
    ):
        values = {
            field: normalize_cell(row[column_index])
            for column_index, field in header_map.items()
        }
        task_name = values.get("task_name", "")
        if not task_name or task_name.startswith("示例："):
            continue

        missing = sorted(field for field in REQUIRED_FIELDS if not values.get(field))
        if missing:
            errors.append(f"第 {row_index} 行缺少必填字段：{', '.join(missing)}")
            continue

        if values["preferred_time_range"] not in {"7d", "14d", "30d", "6m"}:
            errors.append(
                f"第 {row_index} 行首选时间范围必须是 7d、14d、30d 或 6m"
            )
            continue

        for field in MULTI_VALUE_FIELDS:
            values[field] = split_multi_value(values[field])
        briefs.append(values)

    if errors:
        workbook.close()
        raise ValueError("\n".join(errors))
    if not briefs:
        workbook.close()
        raise ValueError("没有找到正式研究任务；请在示例行下方填写一行数据")

    if "keyword_review" in workbook.sheetnames:
        review_sheet = workbook["keyword_review"]
        review_header_row = find_header_row(review_sheet, "关键词")
        review_headers = [normalize_cell(cell.value) for cell in review_sheet[review_header_row]]
        review_header_map = {
            index: REVIEW_FIELDS[header]
            for index, header in enumerate(review_headers)
            if header in REVIEW_FIELDS
        }
        approved_by_task: dict[str, list[dict[str, Any]]] = {}
        review_errors: list[str] = []
        for row_index, row in enumerate(
            review_sheet.iter_rows(
                min_row=review_header_row + 1,
                values_only=True,
            ),
            start=review_header_row + 1,
        ):
            candidate = {
                field: normalize_cell(row[column_index])
                for column_index, field in review_header_map.items()
            }
            task_name = candidate.get("task_name", "")
            keyword = candidate.get("keyword", "")
            if not task_name or task_name.startswith("示例：") or not keyword:
                continue
            approval = candidate.get("approved", "").casefold()
            if approval not in {"", "yes", "no"}:
                review_errors.append(
                    f"keyword_review 第 {row_index} 行的是否采用必须是 yes、no 或留空"
                )
                continue
            if approval != "yes":
                continue
            approved_by_task.setdefault(task_name, []).append(candidate)

        if review_errors:
            workbook.close()
            raise ValueError("\n".join(review_errors))

        for brief in briefs:
            brief["approved_keywords"] = approved_by_task.get(brief["task_name"], [])

    workbook.close()
    return briefs


def sync_keyword_review(workbook_path: Path, briefs: list[dict[str, Any]]) -> int:
    workbook = load_workbook(workbook_path)
    if "keyword_review" not in workbook.sheetnames:
        raise ValueError("缺少 keyword_review 工作表")

    sheet = workbook["keyword_review"]
    header_row = find_header_row(sheet, "关键词")
    headers = [normalize_cell(cell.value) for cell in sheet[header_row]]
    columns = {header: index + 1 for index, header in enumerate(headers)}
    missing_headers = set(REVIEW_FIELDS) - set(headers)
    if missing_headers:
        workbook.close()
        raise ValueError(f"关键词审核表缺少字段：{', '.join(sorted(missing_headers))}")

    existing: dict[tuple[str, str], dict[str, Any]] = {}
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        item = {
            field: normalize_cell(row[index])
            for index, header in enumerate(headers)
            if header in REVIEW_FIELDS
            for field in [REVIEW_FIELDS[header]]
        }
        task_name = item.get("task_name", "")
        keyword = item.get("keyword", "")
        if task_name and keyword and not task_name.startswith("示例："):
            existing[(task_name, keyword.casefold())] = item

    first_generated_row = max(header_row + 1, 12)
    if sheet.max_row >= first_generated_row:
        for row in sheet.iter_rows(
            min_row=first_generated_row,
            max_row=sheet.max_row,
            min_col=1,
            max_col=len(REVIEW_FIELDS),
        ):
            for cell in row:
                cell.value = None

    candidates = [
        candidate
        for brief in briefs
        for candidate in generate_keyword_candidates(brief)
    ]
    output_fields = [REVIEW_FIELDS[header] for header in headers if header in REVIEW_FIELDS]
    for row_index, candidate in enumerate(candidates, start=first_generated_row):
        previous = existing.get((candidate["task_name"], candidate["keyword"].casefold()))
        if previous:
            for field in ("approved", "reason", "priority", "notes"):
                if previous.get(field) != "":
                    candidate[field] = previous[field]
        for field in output_fields:
            header = next(name for name, key in REVIEW_FIELDS.items() if key == field)
            sheet.cell(row=row_index, column=columns[header], value=candidate[field])

    workbook.save(workbook_path)
    workbook.close()
    return len(candidates)


def reviewed_task_names(workbook_path: Path) -> set[str]:
    """Return task names that already have generated rows in keyword_review."""
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if "keyword_review" not in workbook.sheetnames:
            raise ValueError("缺少 keyword_review 工作表")

        sheet = workbook["keyword_review"]
        header_row = find_header_row(sheet, "关键词")
        headers = [normalize_cell(cell.value) for cell in sheet[header_row]]
        columns = {header: index for index, header in enumerate(headers)}
        missing_headers = set(REVIEW_FIELDS) - set(headers)
        if missing_headers:
            raise ValueError(
                f"关键词审核表缺少字段：{', '.join(sorted(missing_headers))}"
            )

        task_column = columns["任务名称"]
        keyword_column = columns["关键词"]
        first_generated_row = max(header_row + 1, 12)
        task_names = set()
        for row in sheet.iter_rows(
            min_row=first_generated_row,
            values_only=True,
        ):
            task_name = normalize_cell(row[task_column])
            keyword = normalize_cell(row[keyword_column])
            if task_name and keyword and not task_name.startswith("示例："):
                task_names.add(task_name)
        return task_names
    finally:
        workbook.close()


def export_briefs_json(briefs: list[dict[str, Any]], output_path: Path) -> Path:
    output_path = output_path.expanduser().resolve()
    output_path.write_text(
        json.dumps({"tasks": briefs}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return output_path


def prepare_research_workbook(
    workbook_path: str | Path,
    output_path: str | Path | None = None,
    *,
    require_review: bool = True,
) -> dict[str, Any]:
    """Validate Excel, synchronize keywords, and export reviewed data to JSON.

    The first submission creates keyword candidates and asks the user to review
    them. A later submission for the same task exports only rows marked yes.
    """
    path = Path(workbook_path).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"找不到输入表：{path}")
    if path.suffix.lower() != ".xlsx":
        raise ValueError("品牌研究输入必须是 .xlsx 文件")

    briefs = read_briefs(path)
    task_names = {brief["task_name"] for brief in briefs}
    already_reviewed = reviewed_task_names(path)
    generated_count = sync_keyword_review(path, briefs)

    if require_review and not task_names.issubset(already_reviewed):
        return {
            "status": "review_required",
            "workbook_path": path,
            "json_path": None,
            "task_count": len(briefs),
            "generated_count": generated_count,
        }

    briefs = read_briefs(path)
    missing_approved = [
        brief["task_name"]
        for brief in briefs
        if not brief.get("approved_keywords")
    ]
    if missing_approved:
        raise ValueError(
            "以下任务没有 approved=yes 的关键词：" + ", ".join(missing_approved)
        )

    json_path = (
        Path(output_path)
        if output_path is not None
        else path.with_name("research_brief.json")
    )
    return {
        "status": "ready",
        "workbook_path": path,
        "json_path": export_briefs_json(briefs, json_path),
        "task_count": len(briefs),
        "generated_count": generated_count,
    }


def main() -> None:
    args = parse_args()
    workbook_path = args.workbook.expanduser().resolve()
    if not workbook_path.exists():
        raise FileNotFoundError(f"找不到输入表：{workbook_path}")

    if args.no_sync_keywords:
        briefs = read_briefs(workbook_path)
        missing_approved = [
            brief["task_name"]
            for brief in briefs
            if not brief.get("approved_keywords")
        ]
        if missing_approved:
            raise ValueError(
                "以下任务没有 approved=yes 的关键词：" + ", ".join(missing_approved)
            )
        output_path = args.output or workbook_path.with_name("research_brief.json")
        output_path = export_briefs_json(briefs, output_path)
        print(f"Validated {len(briefs)} task(s).")
        print(f"JSON written to: {output_path}")
        return

    result = prepare_research_workbook(workbook_path, args.output)
    print(f"Validated {result['task_count']} task(s).")
    print(
        f"Synced {result['generated_count']} keyword candidate(s) to keyword_review."
    )
    if result["status"] == "review_required":
        print("请审核 keyword_review，将采用项设为 yes、不采用项设为 no，然后再次运行。")
    else:
        print(f"JSON written to: {result['json_path']}")


if __name__ == "__main__":
    main()
